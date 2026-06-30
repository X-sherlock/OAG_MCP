from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROJECT_SRC = PROJECT_ROOT / "src"
if str(PROJECT_SRC) not in sys.path:
    sys.path.insert(0, str(PROJECT_SRC))

from oag_mcp.semantic_frame_builder import semantic_frame_for as build_semantic_frame_for

WORKBOOK = Path(os.environ.get("OAG_WORKFLOW_WORKBOOK", PROJECT_ROOT / "OAG基金workflow全量测试问题清单.xlsx"))
OUTPUT_DIR = Path(os.environ.get("OAG_WORKFLOW_OUTPUT_DIR", PROJECT_ROOT / "outputs" / "workflow_full_test"))
CASES_FILE = OUTPUT_DIR / "workflow_cases.json"
PYTHON_OUTPUT = OUTPUT_DIR / "python_agent_plans.json"
JAVA_OUTPUT = OUTPUT_DIR / "java_agent_plans.json"
DIFF_OUTPUT = OUTPUT_DIR / "diff_summary.json"
WORKBOOK_ROWS_OUTPUT = OUTPUT_DIR / "workbook_rows.json"
SUMMARY_OUTPUT = OUTPUT_DIR / "workbook_summary.json"
MAX_EXCEL_CELL_CHARS = 32767


PERIOD_RULES: list[tuple[str, str]] = [
    (r"近(?:一|1)周|一周|最近一周", "1w"),
    (r"近(?:两|2)周|两周|最近两周", "2w"),
    (r"近(?:一|1)个?月|一个月|最近一个月", "1m"),
    (r"近(?:三|3)个?月|三个月|最近三个月", "3m"),
    (r"近(?:半|六|6)年?个?月|半年|最近半年", "6m"),
    (r"近(?:十二|12)个?月|过去12个月", "1y"),
    (r"近(?:一|1)年|一年|过去一年|最近一年", "1y"),
    (r"近(?:两|2)年|两年", "2y"),
    (r"近(?:三|3)年|三年", "3y"),
    (r"近(?:五|5)年|五年", "5y"),
    (r"近(?:十|10)年|十年", "10y"),
    (r"近(?:二十|20)年|二十年", "20y"),
    (r"今年以来|本年以来|\bYTD\b", "ytd"),
    (r"成立以来|设立以来|\bSI\b", "si"),
    (r"最近20个交易日", "20d"),
    (r"最近60个交易日", "60d"),
]

ATTRIBUTE_RULES: list[tuple[str, str]] = [
    ("annualized_return", r"年化收益|年化回报"),
    ("return_rate", r"收益率|收益|回报率|回报|涨幅|涨了|涨得|赚钱|表现"),
    ("benchmark_return", r"基准收益|业绩基准收益|沪深300|中证500|中证偏股基金指数|大盘|市场"),
    ("excess_return", r"超额收益|跑赢|跑输|战胜|优势|强多少|高多少"),
    ("max_drawdown", r"最大回撤|回撤|跌得|下跌|抗跌|跌了|跌幅"),
    ("drawdown", r"连续下跌"),
    ("max_drawdown_peak_date", r"回撤发生|峰值"),
    ("max_drawdown_trough_date", r"回撤发生|谷值"),
    ("max_drawdown_recovery_days", r"修复.*多久|恢复.*多久"),
    ("volatility", r"波动率|波动|稳不稳|更稳|稳定|激进"),
    ("standard_deviation", r"标准差"),
    ("sharpe_ratio", r"夏普|Sharpe|性价比"),
    ("calmar_ratio", r"卡玛|Calmar|收益回撤比"),
    ("tracking_error", r"跟踪误差|走势差异|跟踪偏离"),
    ("information_ratio", r"信息比率|信息比"),
    ("downside_risk", r"下行风险"),
    ("rank", r"排名|排第几|排行|靠前|前10%|前20%"),
    ("percentile", r"分位|百分位|前10%|前20%"),
    ("peer_return_rank", r"同类收益排名|收益排名"),
    ("peer_risk_rank", r"风险控制排名|风险排名|回撤排名"),
    ("peer_sharpe_rank", r"夏普.*排名|夏普.*同类"),
    ("peer_drawdown_rank", r"回撤.*同类|同类.*回撤"),
    ("peer_average", r"同类平均|同类中位数"),
    ("fund_name", r"基金名称|产品名称"),
    ("fund_type", r"基金类型|什么基金类型|分类|品种"),
    ("fund_size", r"基金规模|规模"),
    ("inception_date", r"成立日期|什么时候成立"),
    ("manager_name", r"基金经理|经理是谁"),
    ("company_name", r"基金公司|哪家|管理人"),
    ("benchmark_name", r"业绩比较基准|比较基准|基准是什么|基准名称"),
    ("tracking_index_name", r"跟踪的指数|跟踪指数"),
    ("investment_type", r"投资范围|投资类型"),
    ("stock_name", r"前十大持仓|重仓|股票|持仓"),
    ("stock_nav_ratio", r"持仓占比|股票占比"),
    ("stock_asset_ratio", r"股票仓位|股票资产"),
    ("bond_asset_ratio", r"债券仓位|债券资产"),
    ("cash_asset_ratio", r"现金仓位|现金资产"),
    ("holding_industry", r"行业配置|行业"),
    ("fee_value", r"费率|申购费|赎回费|管理费|托管费"),
    ("fee_type", r"申购费|赎回费|管理费|托管费"),
    ("dividend_per_share", r"分红|派息"),
    ("dividend_date", r"分红.*时候|最近一次分红|分红日期"),
    ("risk_level", r"风险等级|风险偏好|稳健|低风险|风险大"),
    ("report_summary", r"公告|研报"),
    ("news_sentiment", r"新闻"),
]

RELATION_RULES: list[tuple[str, str, str, str]] = [
    ("managed_by", "FundManager", "manager_name", r"基金经理|经理是谁"),
    ("issued_by", "FundCompany", "company_name", r"基金公司|管理人"),
    ("has_benchmark", "Benchmark", "benchmark_name", r"业绩比较基准|比较基准|基准"),
    ("belongs_to_category", "FundCategory", "fund_type", r"基金类型|分类|品种|同类"),
    ("tracks_index", "Index", "tracking_index_name", r"跟踪的指数|跟踪指数"),
    ("has_fee", "FundFee", "fee_value", r"费率|申购费|赎回费|管理费|托管费"),
    ("has_dividend", "Dividend", "dividend_per_share", r"分红|派息"),
    ("has_position", "FundPosition", "stock_name", r"持仓|重仓|股票|债券|行业配置"),
    ("has_asset_allocation", "AssetAllocation", "stock_asset_ratio", r"资产配置|仓位"),
]

INTENT_DEFAULTS: dict[str, dict[str, Any]] = {
    "performance_overview": {
        "task_type": "analyze",
        "attrs": ["return_rate", "benchmark_return", "excess_return", "max_drawdown"],
    },
    "risk_overview": {"task_type": "analyze", "attrs": ["max_drawdown", "volatility"]},
    "benchmark_comparison": {
        "task_type": "compare",
        "attrs": ["return_rate", "benchmark_return", "excess_return"],
    },
    "peer_comparison": {"task_type": "compare", "attrs": ["rank", "percentile", "peer_average"]},
    "fund_profile": {
        "task_type": "profile",
        "attrs": ["fund_name", "fund_type", "manager_name", "company_name", "benchmark_name"],
    },
    "fee_analysis": {"task_type": "query", "attrs": ["fee_value"]},
    "dividend_analysis": {"task_type": "query", "attrs": ["dividend_per_share", "dividend_date"]},
    "holding_analysis": {"task_type": "query", "attrs": ["stock_name", "stock_nav_ratio"]},
    "asset_allocation_analysis": {
        "task_type": "query",
        "attrs": ["stock_asset_ratio", "bond_asset_ratio", "cash_asset_ratio"],
    },
    "fund_comparison": {"task_type": "compare", "attrs": ["return_rate"]},
    "fund_ranking": {"task_type": "rank", "attrs": ["return_rate"]},
    "fund_screening": {"task_type": "screen", "attrs": ["max_drawdown"]},
    "fund_recommendation": {"task_type": "recommend", "attrs": ["return_rate", "max_drawdown"]},
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["generate-cases", "evaluate"])
    args = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if args.command == "generate-cases":
        generate_cases()
    else:
        evaluate_outputs()


def generate_cases() -> None:
    rows = read_questions()
    if len(rows) != 255:
        raise RuntimeError(f"Expected 255 question rows, found {len(rows)}")
    cases = []
    for item in rows:
        frame = build_semantic_frame_for(item["question"], item["category"])
        cases.append(
            {
                "case_id": f"workflow_{int(item['case_no']):03d}",
                "title": f"{item['category']} - {item['question']}",
                "case_no": item["case_no"],
                "category": item["category"],
                "question": item["question"],
                "semantic_frame": frame,
                "user_context": {"permission_scopes": ["fund_public_data:read"]},
            }
        )
    CASES_FILE.write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(cases)} cases to {CASES_FILE}")


def read_questions() -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["测试问题清单"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {name: pos for pos, name in enumerate(headers)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[index["question"]]:
            continue
        rows.append(
            {
                "case_no": int(row[index["case_no"]]),
                "category": str(row[index["category"]]),
                "question": str(row[index["question"]]).strip(),
            }
        )
    return rows


def semantic_frame_for(question: str, category: str) -> dict[str, Any]:
    task_type = infer_task_type(question, category)
    intent = infer_intent(question, category, task_type)
    fund_codes = extract_fund_codes(question)
    attrs = infer_attributes(question, category, intent, task_type)
    constraints = infer_constraints(question, category)
    target_objects = infer_targets(question, category, task_type, fund_codes)
    relation_queries = infer_relation_queries(question, intent)
    ranking = infer_ranking(question, task_type, attrs)
    filters = infer_filters(question, task_type, attrs)
    comparison = infer_comparison(question, task_type, fund_codes, attrs)
    limit = infer_limit(question, task_type)

    if not attrs:
        attrs = list(INTENT_DEFAULTS.get(intent, {}).get("attrs") or [])
    if task_type in {"rank", "screen", "recommend"}:
        mentioned_attributes: list[str] = []
    else:
        mentioned_attributes = unique(attrs)

    frame: dict[str, Any] = {
        "domain": "finance_market",
        "raw_question": question,
        "task_type": task_type,
        "intent": intent,
        "target_objects": target_objects,
        "constraints": constraints,
        "mentioned_attributes": mentioned_attributes,
        "relation_queries": relation_queries,
        "filters": filters,
        "ranking": ranking,
        "comparison": comparison,
        "options": {
            "allow_relation_expansion": True,
            "allow_peer_expansion": True,
            "include_supporting_context": True,
        },
    }
    if limit is not None:
        frame["limit"] = limit
    return frame


def infer_task_type(question: str, category: str) -> str:
    if re.search(r"基金经理|基金公司|基金类型|基金规模|成立日期|跟踪的指数|投资范围|评级|公告|新闻|持仓|重仓|仓位|行业配置|申购费|赎回费|分红|派息", question):
        return "profile" if re.search(r"画像|概况|基本信息", question) else "query"
    if category == "筛选推荐类":
        return "recommend" if re.search(r"推荐|适合|收益高、波动低|稳健", question) else "screen"
    if re.search(r"推荐|适合买吗|稳健型投资|长期持有价值|稳赚不赔", question):
        return "recommend"
    if re.search(r"找出|筛选|选出|过滤", question):
        return "screen"
    if re.search(r"排名前|排行|表现排名|哪个.*最高|前\d+%|前[一二三四五六七八九十]+", question) and not re.search(
        r"同类排名|排名如何|排名是多少|排第几|前十大持仓|前十.*持仓", question
    ):
        return "rank"
    if re.search(r"比较|对比|相比|相对|跑赢|跑输|基准|沪深300|中证500|同类|和\d{6}|与\d{6}|谁的|哪个.*(?:更|最高|好|强)", question):
        return "compare"
    if re.search(r"画像|概况|基本信息", question):
        return "profile"
    if re.search(r"分析|表现|怎么样|如何|好不好|高不高|大不大|稳不稳|行不行", question):
        return "analyze"
    return "query"


def infer_intent(question: str, category: str, task_type: str) -> str:
    if category == "筛选推荐类":
        return "fund_recommendation" if task_type == "recommend" else "fund_screening"
    if task_type == "rank":
        return "fund_ranking"
    if task_type == "screen":
        return "fund_screening"
    if task_type == "recommend":
        return "fund_recommendation"
    if re.search(r"申购费|赎回费|费率|费用", question):
        return "fee_analysis"
    if re.search(r"分红|派息", question):
        return "dividend_analysis"
    if re.search(r"持仓|重仓|行业配置", question):
        return "holding_analysis"
    if re.search(r"仓位|资产配置|股票资产|债券资产|现金资产", question):
        return "asset_allocation_analysis"
    if re.search(r"基金经理|基金公司|基金类型|基金规模|成立日期|跟踪的指数|投资范围|评级|公告|新闻", question):
        return "fund_profile"
    if re.search(r"沪深300|中证500|中证偏股|基准|跑赢|跑输|超额|跟踪误差|信息比率|大盘|市场", question):
        return "benchmark_comparison"
    if re.search(r"同类|排名|分位|中位数", question):
        return "peer_comparison"
    codes = extract_fund_codes(question)
    if task_type == "compare" and len(codes) > 1:
        return "fund_comparison"
    if re.search(r"风险|回撤|波动|夏普|卡玛|下行|稳|激进|低风险", question):
        return "risk_overview"
    return "performance_overview"


def extract_fund_codes(question: str) -> list[str]:
    return unique(re.findall(r"(?<!\d)(\d{6})(?!\d)", question))


def extract_fund_name(question: str) -> str | None:
    quoted = re.search(r"[“\"']([^”\"']{2,40})[”\"']", question)
    if quoted:
        name = quoted.group(1).strip()
        return None if looks_like_non_fund_name(name) else name
    if re.search(r"这只基金|这基金|不存在基金", question):
        return None
    prefix = re.search(
        r"^([\u4e00-\u9fa5A-Za-z0-9]{2,40}?)(?:近|今年|去年|从|收益|最大|风险|和|有没有|跟踪|重仓|股票|债券|行业|分红|费率|成立|基金经理|基金公司|的)",
        question,
    )
    if prefix:
        name = prefix.group(1).strip("，,。？? ")
        if name and not looks_like_non_fund_name(name):
            return name
    return None


def looks_like_non_fund_name(value: str) -> bool:
    if not value or re.fullmatch(r"\d+", value):
        return True
    if re.search(r"帮我|看看|查询|分析|推荐|找出|筛选", value):
        return True
    return bool(re.fullmatch(
        r"(?:近|过去|最近)?(?:一|二|三|四|五|六|七|八|九|十|十二|1|2|3|4|5|6|7|8|9|10|12)?(?:天|周|个月|月|年|季度)|今年以来|成立以来|YTD|SI",
        value,
        flags=re.IGNORECASE,
    ))


def infer_targets(question: str, category: str, task_type: str, fund_codes: list[str]) -> list[dict[str, Any]]:
    if not fund_codes and re.search(r"这只基金|这基金", question):
        role = "comparison_subject" if task_type == "compare" else "analysis_subject"
        return [{"object_type": "Fund", "instance_ref": {}, "role": role}]
    if task_type in {"rank", "screen", "recommend"}:
        return [
            {
                "object_type": "FundSet",
                "instance_ref": {"fund_universe": fund_universe(question)},
                "role": "candidate_set",
            }
        ]
    role = "comparison_subject" if task_type == "compare" and len(fund_codes) > 1 else "analysis_subject"
    if fund_codes:
        return [
            {"object_type": "Fund", "instance_ref": {"fund_code": code}, "role": role}
            for code in fund_codes
        ]
    fund_name = extract_fund_name(question)
    if fund_name:
        return [{"object_type": "Fund", "instance_ref": {"fund_name": fund_name}, "role": role}]
    if re.search(r"这只基金|这基金|基金", question):
        return [{"object_type": "Fund", "instance_ref": {}, "role": role}]
    return [{"object_type": "Fund", "instance_ref": {}, "role": role}]


def fund_universe(question: str) -> str:
    if re.search(r"主动权益", question):
        return "active_equity_funds"
    if re.search(r"权益|股票", question):
        return "equity_funds"
    if re.search(r"债券|债基", question):
        return "bond_funds"
    if re.search(r"指数|ETF", question):
        return "index_funds"
    if re.search(r"混合", question):
        return "mixed_funds"
    return "all_public_funds"


def infer_constraints(question: str, category: str) -> dict[str, Any]:
    constraints: dict[str, Any] = {}
    for pattern, code in PERIOD_RULES:
        if re.search(pattern, question, re.IGNORECASE):
            constraints["period"] = code
            break
    range_match = re.search(
        r"从(20\d{2})[-年/.](\d{1,2})[-月/.](\d{1,2})日?到(20\d{2})[-年/.](\d{1,2})[-月/.](\d{1,2})日?",
        question,
    )
    if range_match:
        y1, m1, d1, y2, m2, d2 = map(int, range_match.groups())
        constraints["date_range"] = {
            "start_date": f"{y1:04d}-{m1:02d}-{d1:02d}",
            "end_date": f"{y2:04d}-{m2:02d}-{d2:02d}",
        }
        constraints.setdefault("period", "custom")
    if "2025年上半年" in question:
        constraints["date_range"] = {"start_date": "2025-01-01", "end_date": "2025-06-30"}
        constraints.setdefault("period", "custom")
    if "去年全年" in question:
        constraints["date_range"] = {"relative": "last_calendar_year"}
        constraints.setdefault("period", "last_year")
    if re.search(r"最新|当前|最近一次", question) or re.search(r"持仓|仓位|资产配置", question):
        constraints.setdefault("report_date", "latest")
    if re.search(r"明天|未来|2099|近100年|必须保证|稳赚不赔|一定会涨|预测明天", question):
        constraints["test_flag"] = "invalid_or_future_or_compliance_risk"
    return constraints


def infer_attributes(question: str, category: str, intent: str, task_type: str) -> list[str]:
    attrs: list[str] = []
    for name, pattern in ATTRIBUTE_RULES:
        if re.search(pattern, question, re.IGNORECASE):
            attrs.append(name)
    if re.search(r"股票仓位|债券仓位|现金仓位|资产配置|股票资产|债券资产|现金资产", question):
        attrs = [item for item in attrs if item not in {"stock_name", "stock_nav_ratio", "bond_name", "bond_nav_ratio", "holding_industry"}]
    if "风险收益" in question:
        attrs.extend(["return_rate", "max_drawdown", "volatility", "sharpe_ratio"])
    if "高收益低回撤" in question or "收益高、波动低" in question:
        attrs.extend(["return_rate", "max_drawdown", "volatility"])
    if "谁表现更好" in question or "哪个收益" in question:
        attrs.append("return_rate")
    if "哪个更稳" in question:
        attrs.extend(["max_drawdown", "volatility"])
    if "适合低风险" in question or "稳健" in question:
        attrs.extend(["max_drawdown", "volatility", "risk_level"])
    if "同类" in question and not any(a in attrs for a in ["rank", "percentile", "peer_average"]):
        attrs.extend(["rank", "percentile", "peer_average"])
    if "基准" in question or "沪深300" in question or "中证500" in question:
        attrs.extend(["benchmark_return", "excess_return"])
    if task_type in {"rank", "screen", "recommend"} and not attrs:
        attrs.extend(INTENT_DEFAULTS.get(intent, {}).get("attrs") or ["return_rate"])
    return unique(attrs)


def infer_relation_queries(question: str, intent: str) -> list[dict[str, Any]]:
    rows = []
    for relation_type, target_type, attribute, pattern in RELATION_RULES:
        if re.search(pattern, question):
            if relation_type == "has_position" and re.search(r"股票仓位|债券仓位|现金仓位|资产配置|股票资产|债券资产|现金资产", question):
                continue
            rows.append(
                {
                    "relation_type": relation_type,
                    "target_object_type": target_type,
                    "attribute_name": attribute,
                }
            )
    if intent == "peer_comparison" and not any(row["relation_type"] == "belongs_to_category" for row in rows):
        rows.append(
            {
                "relation_type": "belongs_to_category",
                "target_object_type": "FundCategory",
                "attribute_name": "fund_type",
            }
        )
    if intent == "benchmark_comparison" and not any(row["relation_type"] == "has_benchmark" for row in rows):
        rows.append(
            {
                "relation_type": "has_benchmark",
                "target_object_type": "Benchmark",
                "attribute_name": "benchmark_name",
            }
        )
    return rows


def infer_ranking(question: str, task_type: str, attrs: list[str]) -> list[dict[str, Any]]:
    if task_type not in {"rank", "recommend"}:
        return []
    rows = []
    if re.search(r"收益|表现|年化", question):
        rows.append({"attribute": "return_rate", "direction": "desc"})
    if re.search(r"回撤.*小|回撤.*低|低回撤|回撤较小|后50%", question):
        rows.append({"attribute": "max_drawdown", "direction": "asc"})
    if re.search(r"波动低|稳健|稳定", question):
        rows.append({"attribute": "volatility", "direction": "asc"})
    if re.search(r"夏普", question):
        rows.append({"attribute": "sharpe_ratio", "direction": "desc"})
    if re.search(r"卡玛", question):
        rows.append({"attribute": "calmar_ratio", "direction": "desc"})
    if re.search(r"信息比率", question):
        rows.append({"attribute": "information_ratio", "direction": "desc"})
    if not rows and attrs:
        rows.append({"attribute": attrs[0], "direction": "desc"})
    return rows or [{"attribute": "return_rate", "direction": "desc"}]


def infer_filters(question: str, task_type: str, attrs: list[str]) -> list[dict[str, Any]]:
    if task_type not in {"screen", "recommend"}:
        return []
    rows = []
    for attr, alias in [
        ("max_drawdown", r"最大回撤|回撤"),
        ("return_rate", r"收益率|年化收益率|收益"),
        ("volatility", r"波动"),
    ]:
        pattern = rf"(?:{alias}).*?(小于|低于|不超过|大于|高于|超过)\s*(\d+(?:\.\d+)?)\s*%"
        match = re.search(pattern, question)
        if match:
            op_text, value = match.groups()
            operator = "<=" if op_text in {"小于", "低于", "不超过"} else ">="
            rows.append({"attribute": attr, "operator": operator, "value": float(value) / 100.0})
    if "收益率大于20%" in question:
        rows.append({"attribute": "return_rate", "operator": ">=", "value": 0.2})
    if "年化收益率大于15%" in question:
        rows.append({"attribute": "annualized_return", "operator": ">=", "value": 0.15})
    if "前10%" in question:
        rows.append({"attribute": "percentile", "operator": "<=", "value": 0.1})
    if "前20%" in question:
        rows.append({"attribute": "percentile", "operator": "<=", "value": 0.2})
    if "后50%" in question:
        rows.append({"attribute": "peer_drawdown_rank", "operator": ">=", "value": 0.5})
    if "回撤较小" in question or "低回撤" in question:
        rows.append({"attribute": "max_drawdown", "operator": "<=", "value": 0.1})
    if "波动低" in question:
        rows.append({"attribute": "volatility", "operator": "<=", "value": 0.15})
    return unique_dicts(rows)


def infer_comparison(question: str, task_type: str, fund_codes: list[str], attrs: list[str]) -> dict[str, Any]:
    if task_type != "compare":
        return {}
    comparison: dict[str, Any] = {
        "mode": "side_by_side",
        "attributes": unique(attrs or ["return_rate"]),
        "target_object_policy": "all_targets",
    }
    if re.search(r"沪深300|中证500|中证偏股基金指数|大盘|市场|基准", question):
        comparison["benchmark_refs"] = benchmark_refs(question)
    if re.search(r"同类|中位数|平均", question):
        comparison["peer_context"] = True
    return comparison


def benchmark_refs(question: str) -> list[dict[str, str]]:
    refs = []
    for name in ["沪深300", "中证500", "中证偏股基金指数"]:
        if name in question:
            refs.append({"benchmark_name": name})
    if not refs and re.search(r"基准|大盘|市场", question):
        refs.append({"benchmark_name": "semantic_or_default_benchmark"})
    return refs


def infer_limit(question: str, task_type: str) -> int | None:
    if task_type not in {"rank", "screen", "recommend"}:
        return None
    match = re.search(r"前\s*(\d+)", question, re.IGNORECASE)
    if match:
        return int(match.group(1))
    match = re.search(r"前([一二三四五六七八九十]+)", question)
    if match:
        return chinese_number(match.group(1))
    return 10 if task_type == "recommend" else 20


def chinese_number(value: str) -> int:
    mapping = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
    if value == "十":
        return 10
    if value.startswith("十"):
        return 10 + mapping.get(value[1:], 0)
    if "十" in value:
        tens, ones = value.split("十", 1)
        return mapping.get(tens, 1) * 10 + mapping.get(ones, 0)
    return mapping.get(value, 10)


def evaluate_outputs() -> None:
    cases = json.loads(CASES_FILE.read_text(encoding="utf-8"))
    python_rows = load_output(PYTHON_OUTPUT)
    java_rows = load_output(JAVA_OUTPUT)
    assert_alignment(cases, python_rows, java_rows)
    tested_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    workbook_rows = []
    diff_rows = []
    summary_counter: dict[str, Counter[str]] = defaultdict(Counter)
    outcome_counter: Counter[str] = Counter()
    category_outcome_counter: dict[str, Counter[str]] = defaultdict(Counter)
    failure_cause_counter: Counter[str] = Counter()
    category_failure_cause_counter: dict[str, Counter[str]] = defaultdict(Counter)
    issue_counter: Counter[str] = Counter()
    exact_match_count = 0

    for case, py_row, java_row in zip(cases, python_rows, java_rows):
        py_plan = py_row.get("agent_plan") or {}
        java_plan = java_row.get("agent_plan") or {}
        is_exact = py_plan == java_plan
        exact_match_count += int(is_exact)
        diff = diff_plan(py_plan, java_plan)
        (
            grade,
            score,
            evaluation,
            issues,
            evaluation_outcome,
            expected_terminal_behavior,
            true_planner_error,
            failure_cause,
            oag_reasoning_error,
            evaluator_reason_zh,
        ) = evaluate_case(case, py_plan, java_plan, diff)
        summary_counter[case["category"]][grade] += 1
        outcome_counter[evaluation_outcome] += 1
        category_outcome_counter[case["category"]][evaluation_outcome] += 1
        failure_cause_counter[failure_cause] += 1
        category_failure_cause_counter[case["category"]][failure_cause] += 1
        for issue in issues:
            issue_counter[issue] += 1
        diff_text = "一致" if is_exact else ("；".join(diff)[:3000] if diff else "完整 JSON 不完全一致，但核心规划字段一致")
        diff_rows.append(
            {
                "case_no": case["case_no"],
                "case_id": case["case_id"],
                "category": case["category"],
                "exact_match": is_exact,
                "diff": diff,
                "grade": grade,
                "score": score,
                "issues": issues,
                "evaluation_outcome": evaluation_outcome,
                "expected_terminal_behavior": expected_terminal_behavior,
                "true_planner_error": true_planner_error,
                "failure_cause": failure_cause,
                "oag_reasoning_error": oag_reasoning_error,
            }
        )
        workbook_rows.append(
            {
                "case_no": case["case_no"],
                "semantic_frame_json": cell_json(case["semantic_frame"], f"{case['case_id']}_semantic_frame.json"),
                "python_agent_plan_json": cell_json(py_plan, f"{case['case_id']}_python_agent_plan.json"),
                "java_agent_plan_json": cell_json(java_plan, f"{case['case_id']}_java_agent_plan.json"),
                "python_summary": summarize_plan(py_plan),
                "java_summary": summarize_plan(java_plan),
                "java_python_diff": diff_text,
                "evaluation_grade": grade,
                "evaluation_score": score,
                "evaluation_zh": evaluation,
                "issues_zh": "；".join(issues) if issues else "",
                "evaluation_outcome": evaluation_outcome,
                "expected_terminal_behavior": expected_terminal_behavior,
                "true_planner_error": "是" if true_planner_error else "否",
                "failure_cause": failure_cause,
                "oag_reasoning_error": "是" if oag_reasoning_error else "否",
                "evaluator_reason_zh": evaluator_reason_zh,
                "tested_at": tested_at,
            }
        )

    summary = build_summary(
        cases,
        summary_counter,
        issue_counter,
        exact_match_count,
        outcome_counter,
        category_outcome_counter,
        failure_cause_counter,
        category_failure_cause_counter,
    )
    DIFF_OUTPUT.write_text(json.dumps(diff_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    WORKBOOK_ROWS_OUTPUT.write_text(json.dumps(workbook_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    SUMMARY_OUTPUT.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote evaluation for {len(workbook_rows)} rows")


def load_output(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def assert_alignment(cases: list[dict[str, Any]], python_rows: list[dict[str, Any]], java_rows: list[dict[str, Any]]) -> None:
    if not (len(cases) == len(python_rows) == len(java_rows) == 255):
        raise RuntimeError(f"Row counts mismatch: cases={len(cases)}, python={len(python_rows)}, java={len(java_rows)}")
    case_ids = [item["case_id"] for item in cases]
    if [item["case_id"] for item in python_rows] != case_ids:
        raise RuntimeError("Python output case_id order does not match cases")
    if [item["case_id"] for item in java_rows] != case_ids:
        raise RuntimeError("Java output case_id order does not match cases")


def diff_plan(left: dict[str, Any], right: dict[str, Any]) -> list[str]:
    diffs = []
    for path, getter in [
        ("status", lambda x: x.get("status")),
        ("task", lambda x: x.get("task")),
        ("coverage.coverage_status", lambda x: (x.get("coverage") or {}).get("coverage_status")),
        ("execution.execution_status", lambda x: (x.get("execution") or {}).get("execution_status")),
        ("facts", lambda x: fact_signature(x)),
        ("skill_calls", lambda x: skill_signature(x)),
        ("issues", lambda x: issue_signature(x)),
    ]:
        lv = getter(left)
        rv = getter(right)
        if lv != rv:
            diffs.append(f"{path} 不一致：Python={compact(lv)} / Java={compact(rv)}")
    return diffs


def fact_signature(plan: dict[str, Any]) -> list[tuple[Any, ...]]:
    return [
        (
            item.get("fact_type"),
            item.get("attribute"),
            item.get("priority"),
            item.get("source"),
            item.get("target_object_type"),
        )
        for item in plan.get("facts") or []
    ]


def skill_signature(plan: dict[str, Any]) -> list[tuple[Any, ...]]:
    return [
        (
            item.get("skill_id"),
            tuple(sorted(item.get("missing_params") or [])),
            item.get("call_status"),
            item.get("covers_required_count"),
            item.get("covers_optional_count"),
        )
        for item in plan.get("skill_calls") or []
    ]


def issue_signature(plan: dict[str, Any]) -> list[tuple[Any, ...]]:
    return [(item.get("code"), item.get("severity"), item.get("message_zh")) for item in plan.get("issues") or []]


def evaluate_case(
    case: dict[str, Any], py_plan: dict[str, Any], java_plan: dict[str, Any], diff: list[str]
) -> tuple[str, int, str, list[str], str, str, bool, str, bool, str]:
    frame = case["semantic_frame"]
    question = case["question"]
    category = case["category"]
    issues: list[str] = []
    plan = py_plan if py_plan else java_plan
    if not py_plan or not java_plan:
        return (
            "ERR",
            0,
            "OAG 调用未产生有效 agent_plan。",
            ["调用异常或输出缺失"],
            "true_planner_error",
            "ready_answer",
            True,
            "oag_reasoning_error",
            True,
            "Python 或 Java 未返回有效 agent_plan。",
        )
    status = plan.get("status")
    coverage = plan.get("coverage") or {}
    execution = plan.get("execution") or {}
    skill_calls = plan.get("skill_calls") or []
    facts = plan.get("facts") or []
    fact_attrs = {item.get("attribute") for item in facts if item.get("attribute")}
    required_facts = [item for item in facts if item.get("priority") == "required"]
    required_attrs = {item.get("attribute") for item in required_facts if item.get("attribute")}
    expected_attrs = expected_attributes(frame)
    missing_expected = sorted(expected_attrs - fact_attrs)
    allowed_attrs = allowed_supporting_attributes(question, frame)
    extra_required_attrs = sorted(required_attrs - expected_attrs - allowed_attrs)
    relation_issues = relation_quality_issues(frame, plan)
    false_compare_issue = has_false_compare_target_issue(frame, plan)
    uncovered_required = coverage.get("uncovered_required_facts") or []
    terminal = classify_terminal_behavior(case, plan, diff)
    evaluation_outcome = terminal["evaluation_outcome"]
    expected_terminal_behavior = terminal["expected_terminal_behavior"]
    correct_terminal_behavior = bool(terminal["correct_terminal_behavior"])
    true_planner_error = bool(terminal["true_planner_error"])
    failure_cause = str(terminal["failure_cause"])
    oag_reasoning_error = bool(terminal["oag_reasoning_error"])
    evaluator_reason_zh = terminal["evaluator_reason_zh"]
    answerability_issues = []
    redundancy_issues = []
    abnormal_issues = []

    if status == "error":
        answerability_issues.append("规划返回 error，无法形成可执行回答链路")
    if not correct_terminal_behavior:
        if category == "缺基金参数":
            if not has_missing(skill_calls, "fund_code") and not has_issue_code(plan, {"FUND_CODE_REQUIRED", "INVALID_FUND_CODE_FORMAT"}):
                abnormal_issues.append("缺基金参数未形成 fund_code 阻塞，可能会在无基金对象时继续规划")
        elif category == "缺时间参数":
            if not has_missing(skill_calls, "period") and not has_unsupported_or_clarification_status(plan):
                abnormal_issues.append("缺时间参数未形成 period 阻塞或不支持提示，可能会回答无时间范围的指标")
        elif is_invalid_or_compliance(question):
            if not has_rejection_or_risk_issue(plan):
                abnormal_issues.append("非法/未来/合规风险输入没有被 OAG 显式拒绝或标记")
            if execution.get("execution_status") == "ready":
                abnormal_issues.append("非法/未来/合规风险输入仍被规划为可直接执行")

    if not correct_terminal_behavior:
        if coverage.get("coverage_status") not in {"full_coverage", "partial_coverage"}:
            answerability_issues.append(f"覆盖状态为 {coverage.get('coverage_status')}，无法确认能回答用户问题")
        if execution.get("execution_status") in {"no_skill_calls", "disabled"}:
            answerability_issues.append(f"执行状态为 {execution.get('execution_status')}，没有可用 Skill 调用")
        if execution.get("execution_status") == "blocked_missing_params":
            answerability_issues.append("普通问题出现缺参阻塞，当前输出不能直接回答用户")
        if missing_expected:
            answerability_issues.append("用户问题所需关键属性未进入事实需求：" + ",".join(missing_expected[:8]))

    if uncovered_required and not correct_terminal_behavior:
        answerability_issues.append(f"存在 {len(uncovered_required)} 个必需事实未被 Skill 覆盖")
    if relation_issues and not correct_terminal_behavior:
        issues.extend(relation_issues)
    if false_compare_issue:
        relation_issues.append("基准/同类比较被误报 COMPARE_TARGET_TOO_FEW；单基金对基准/同类关系不应要求两个 Fund")
    if extra_required_attrs and not correct_terminal_behavior:
        redundancy_issues.append("输出了用户未要求的必需事实：" + ",".join(extra_required_attrs[:10]))
    extra_skills = unnecessary_skill_calls(frame, plan, expected_attrs | allowed_attrs)
    if extra_skills and not correct_terminal_behavior:
        redundancy_issues.append("存在疑似不必要 Skill 调用：" + ",".join(extra_skills[:8]))
    if diff:
        issues.append("Java/Python 输出不一致，需要避免两版关系扩展或 Skill 覆盖规则漂移")
    if oag_reasoning_error:
        issues.append("真实 planner error：" + evaluator_reason_zh)

    issues = unique(answerability_issues + relation_issues + redundancy_issues + abnormal_issues + issues)
    if correct_terminal_behavior and evaluation_outcome != "answerable_ready":
        score = {
            "correct_blocked_missing_params": 100,
            "correct_unsupported_data_source": 100,
            "correct_unsupported_period": 100,
            "correct_rejected_guardrail": 100,
        }.get(evaluation_outcome, 88)
        score -= 10 if diff else 0
    else:
        score = 100
        score -= 40 if status == "error" else 0
        score -= min(35, len(uncovered_required) * 15)
        score -= min(30, len(missing_expected) * 12)
        score -= min(35, len(relation_issues) * 15)
        score -= min(25, len(extra_required_attrs) * 5)
        score -= min(15, len(extra_skills) * 5)
        score -= 25 if execution.get("execution_status") in {"no_skill_calls", "disabled"} else 0
        score -= 20 if execution.get("execution_status") == "blocked_missing_params" and category not in {"缺基金参数", "缺时间参数", "非法异常输入"} else 0
        score -= 10 if diff else 0
        if category in {"缺基金参数", "缺时间参数"}:
            if not abnormal_issues:
                score = min(score, 92)
            else:
                score -= 20
        if is_invalid_or_compliance(question):
            score = min(score, 75)
            if abnormal_issues:
                score -= 20
    score = max(0, score)
    if (uncovered_required and not correct_terminal_behavior) or relation_issues:
        score = min(score, 74)
    if extra_required_attrs and len(extra_required_attrs) >= 3:
        score = min(score, 79)
    if false_compare_issue:
        score = min(score, 69)
    if answerability_issues and execution.get("execution_status") in {"no_skill_calls", "disabled"}:
        score = min(score, 54)
    if oag_reasoning_error:
        score = min(score, 54)
    if score >= 90:
        grade = "A"
    elif score >= 75:
        grade = "B"
    elif score >= 55:
        grade = "C"
    else:
        grade = "D"
    if status == "error" and score < 40:
        grade = "D"
    evaluation = evaluation_text(
        grade,
        coverage,
        execution,
        len(facts),
        len(skill_calls),
        bool(diff),
        answerability_issues,
        relation_issues,
        redundancy_issues,
        abnormal_issues,
        evaluation_outcome,
        expected_terminal_behavior,
        true_planner_error,
        failure_cause,
        oag_reasoning_error,
        evaluator_reason_zh,
    )
    return (
        grade,
        score,
        evaluation,
        issues,
        evaluation_outcome,
        expected_terminal_behavior,
        true_planner_error,
        failure_cause,
        oag_reasoning_error,
        evaluator_reason_zh,
    )


def expected_attributes(frame: dict[str, Any]) -> set[str]:
    attrs = set(frame.get("mentioned_attributes") or [])
    for item in frame.get("ranking") or []:
        if item.get("attribute"):
            attrs.add(item["attribute"])
    for item in frame.get("filters") or []:
        if item.get("attribute"):
            attrs.add(item["attribute"])
    for item in frame.get("comparison", {}).get("attributes") or []:
        attrs.add(item)
    for item in frame.get("relation_queries") or []:
        if item.get("attribute_name"):
            attrs.add(item["attribute_name"])
    return {item for item in attrs if item not in {"fee_type"}}


def allowed_supporting_attributes(question: str, frame: dict[str, Any]) -> set[str]:
    allowed: set[str] = set()
    intent = str(frame.get("intent") or "")
    if re.search(r"表现|分析|怎么样|好不好|风险收益|高不高|大不大", question):
        allowed.update({"return_rate", "max_drawdown", "volatility", "sharpe_ratio"})
    if intent == "performance_overview":
        allowed.update({"benchmark_return", "excess_return"})
    if intent == "risk_overview":
        allowed.update({"return_rate", "max_drawdown", "volatility", "sharpe_ratio", "calmar_ratio", "downside_risk"})
    if intent == "benchmark_comparison" or re.search(r"基准|沪深300|中证500|大盘|市场|跑赢|跑输|超额", question):
        allowed.update({"benchmark_return", "excess_return", "tracking_error", "information_ratio", "benchmark_name"})
    if intent == "peer_comparison" or re.search(r"同类|排名|分位|中位数|平均", question):
        allowed.update({"rank", "percentile", "peer_average", "peer_return_rank", "peer_risk_rank", "peer_sharpe_rank", "peer_drawdown_rank", "fund_type"})
    if intent == "fund_profile" and re.search(r"画像|概况|基本信息", question):
        allowed.update({"fund_name", "fund_type", "manager_name", "company_name", "benchmark_name", "risk_level", "fund_size", "inception_date"})
    return allowed


def relation_quality_issues(frame: dict[str, Any], plan: dict[str, Any]) -> list[str]:
    issues = []
    relation_queries = frame.get("relation_queries") or []
    if not relation_queries:
        return issues
    facts = plan.get("facts") or []
    plan_issues = plan.get("issues") or []
    relation_attrs = {item.get("attribute") for item in facts if item.get("fact_type") == "relation_instance"}
    uncovered_relation = any((item.get("code") or "").startswith("RELATION_INSTANCE_UNCOVERED") for item in plan_issues)
    for query in relation_queries:
        attr = query.get("attribute_name")
        relation_type = query.get("relation_type")
        if attr and attr not in relation_attrs:
            issues.append(f"关系 {relation_type} 未生成对应 relation_instance 事实（属性 {attr}）")
    if uncovered_relation:
        issues.append("关系事实已生成但未被 Skill 覆盖，无法支撑最终回答")
    return unique(issues)


def has_false_compare_target_issue(frame: dict[str, Any], plan: dict[str, Any]) -> bool:
    if frame.get("intent") not in {"benchmark_comparison", "peer_comparison"}:
        return False
    if frame.get("task_type") != "compare":
        return False
    return any(item.get("code") == "COMPARE_TARGET_TOO_FEW" for item in plan.get("issues") or [])


def classify_terminal_behavior(case: dict[str, Any], plan: dict[str, Any], diff: list[str]) -> dict[str, Any]:
    frame = case["semantic_frame"]
    question = case["question"]
    expected = expected_terminal_behavior_for(case, frame, question)
    status = str(plan.get("status") or "")
    coverage_status = str((plan.get("coverage") or {}).get("coverage_status") or "")
    execution_status = str((plan.get("execution") or {}).get("execution_status") or "")
    codes = issue_codes(plan)

    actual = "answerable_ready"
    reason = "planner 生成可执行规划。"
    if has_guardrail_rejection(plan):
        actual = "correct_rejected_guardrail"
        reason = "planner 对未来收益、保收益或合规风险请求给出了显式拒绝。"
    elif has_fund_param_block(plan):
        actual = "correct_blocked_missing_params"
        reason = "planner 识别到缺少可执行所需的有效基金对象或基金代码。"
    elif has_unsupported_or_clarification_status(plan):
        if is_unsupported_period_question(question, frame) or "UNSUPPORTED_LONG_PERIOD" in codes or "INVALID_DATE_RANGE" in codes:
            actual = "correct_unsupported_period"
            reason = "planner 识别到当前周期、区间或日期条件不在 ifund_all_info 支持范围内。"
        elif has_unsupported_data_issue(plan):
            actual = "correct_unsupported_data_source"
            reason = "planner 明确声明请求的数据或关系超出 ifund_all_info 当前可用能力。"
        else:
            actual = "ambiguous_needs_review"
            reason = "planner 进入澄清或不可执行终态，但缺少足够明确的阻塞类型。"
    elif status == "success" and coverage_status in {"full_coverage", "partial_coverage"} and execution_status == "ready":
        actual = "answerable_ready"
        reason = "planner 生成 ready 状态的可执行规划。"
    else:
        actual = "true_planner_error"
        reason = f"planner 输出 {status}/{coverage_status}/{execution_status}，既不可执行，也不是显式正确阻塞、拒绝或 unsupported。"

    expected_to_actual = {
        "ready_answer": "answerable_ready",
        "blocked_missing_params": "correct_blocked_missing_params",
        "unsupported_data_source": "correct_unsupported_data_source",
        "unsupported_period": "correct_unsupported_period",
        "rejected_guardrail": "correct_rejected_guardrail",
    }
    correct_terminal_behavior = actual == expected_to_actual.get(expected)

    reported_outcome = actual
    true_planner_error = False
    failure_cause = failure_cause_for_actual(actual)
    oag_reasoning_error = False
    if diff:
        true_planner_error = True
        reported_outcome = "true_planner_error"
        failure_cause = "oag_reasoning_error"
        oag_reasoning_error = True
        reason = "Python / Java agent_plan 不一致，属于跨语言规划链路错误。"
    elif actual == "true_planner_error":
        true_planner_error = True
        reported_outcome = "true_planner_error"
        failure_cause = "oag_reasoning_error"
        oag_reasoning_error = True
    elif actual == "ambiguous_needs_review":
        true_planner_error = False
        failure_cause = "needs_review"
    elif not correct_terminal_behavior:
        true_planner_error = True
        reported_outcome = "true_planner_error"
        failure_cause = "oag_reasoning_error"
        oag_reasoning_error = True
        reason = f"期望终态为 {expected}，但 planner 实际终态为 {actual}。"

    return {
        "evaluation_outcome": reported_outcome,
        "expected_terminal_behavior": expected,
        "correct_terminal_behavior": correct_terminal_behavior,
        "true_planner_error": true_planner_error,
        "failure_cause": failure_cause,
        "oag_reasoning_error": oag_reasoning_error,
        "evaluator_reason_zh": reason,
    }


def failure_cause_for_actual(actual: str) -> str:
    if actual == "answerable_ready":
        return "none"
    if actual == "correct_blocked_missing_params":
        return "input_missing"
    if actual in {"correct_unsupported_data_source", "correct_unsupported_period"}:
        return "system_limitation"
    if actual == "correct_rejected_guardrail":
        return "guardrail_policy"
    if actual == "ambiguous_needs_review":
        return "needs_review"
    return "oag_reasoning_error"


def expected_terminal_behavior_for(case: dict[str, Any], frame: dict[str, Any], question: str) -> str:
    if is_guardrail_question(question):
        return "rejected_guardrail"
    if is_missing_required_param_question(case, frame, question):
        return "blocked_missing_params"
    if is_unsupported_period_question(question, frame):
        return "unsupported_period"
    if has_unsupported_data_request(frame, question):
        return "unsupported_data_source"
    return "ready_answer"


def is_missing_required_param_question(case: dict[str, Any], frame: dict[str, Any], question: str) -> bool:
    return is_missing_fund_question(case, frame, question)


def is_missing_fund_question(case: dict[str, Any], frame: dict[str, Any], question: str) -> bool:
    if case.get("category") == "缺基金参数":
        return True
    if not extract_fund_codes(question) and re.search(r"这只基金|这基金|不存在基金|和沪深300比", question):
        return True
    return any(
        item.get("object_type") == "Fund" and not has_fund_identifier(item.get("instance_ref") or {})
        for item in frame.get("target_objects") or []
    ) and case.get("category") in {"实体识别", "非法异常输入"}


def has_fund_identifier(instance_ref: dict[str, Any]) -> bool:
    return bool(instance_ref.get("fund_code") or instance_ref.get("fund_name") or instance_ref.get("fund_short_name"))


def is_guardrail_question(question: str) -> bool:
    return bool(re.search(r"明天|未来|2099|必须保证|稳赚不赔|一定会涨|预测明天", question))


def is_unsupported_period_question(question: str, frame: dict[str, Any]) -> bool:
    constraints = frame.get("constraints") or {}
    period = constraints.get("period")
    if period in {"2w", "10y", "20y", "si", "custom", "last_year"}:
        return True
    if constraints.get("date_range"):
        return True
    return bool(re.search(r"近(?:两|2)周|近(?:十|10)年|近(?:二十|20)年|近100年|成立以来|最近(?:20|60)个交易日|去年全年|上半年", question))


def has_unsupported_data_request(frame: dict[str, Any], question: str) -> bool:
    attrs = expected_attributes(frame)
    unsupported_attrs = {
        "tracking_error",
        "information_ratio",
        "tracking_index_name",
        "investment_type",
        "stock_name",
        "stock_nav_ratio",
        "holding_industry",
        "fee_value",
        "dividend_per_share",
        "dividend_date",
        "report_summary",
        "news_sentiment",
    }
    if attrs.intersection(unsupported_attrs):
        return True
    return bool(re.search(r"评级|公告|新闻|前十大持仓|重仓|行业配置|分红|跟踪的指数|投资范围|费率", question))


def issue_codes(plan: dict[str, Any]) -> set[str]:
    return {str(item.get("code")) for item in plan.get("issues") or [] if item.get("code")}


def has_issue_code(plan: dict[str, Any], codes: set[str]) -> bool:
    return bool(issue_codes(plan).intersection(codes))


def has_guardrail_rejection(plan: dict[str, Any]) -> bool:
    status = str(plan.get("status") or "")
    coverage_status = str((plan.get("coverage") or {}).get("coverage_status") or "")
    codes = issue_codes(plan)
    return (
        status in {"rejected", "rejected_with_risk_notice"}
        or coverage_status in {"rejected", "rejected_with_risk_notice"}
        or bool(codes.intersection({"FUTURE_RETURN_QUERY_REJECTED", "FUTURE_DATE_RANGE_REJECTED", "GUARANTEED_PROFIT_REJECTED"}))
    )


def has_fund_param_block(plan: dict[str, Any]) -> bool:
    status = str(plan.get("status") or "")
    coverage_status = str((plan.get("coverage") or {}).get("coverage_status") or "")
    execution_status = str((plan.get("execution") or {}).get("execution_status") or "")
    return (
        status == "blocked_missing_params"
        or coverage_status == "blocked_missing_params"
        or execution_status == "blocked_missing_params"
        or has_issue_code(plan, {"FUND_CODE_REQUIRED", "INVALID_FUND_CODE_FORMAT"})
    )


def has_unsupported_or_clarification_status(plan: dict[str, Any]) -> bool:
    status = str(plan.get("status") or "")
    coverage_status = str((plan.get("coverage") or {}).get("coverage_status") or "")
    execution_status = str((plan.get("execution") or {}).get("execution_status") or "")
    return (
        status in {"unsupported", "need_clarification"}
        or coverage_status in {"unsupported", "need_clarification"}
        or execution_status in {"unsupported", "need_clarification"}
        or has_issue_code(plan, {"UNSUPPORTED_DATA_CAPABILITY", "UNSUPPORTED_LONG_PERIOD", "INVALID_DATE_RANGE"})
    )


def has_unsupported_data_issue(plan: dict[str, Any]) -> bool:
    return has_issue_code(plan, {"UNSUPPORTED_DATA_CAPABILITY"})


def unnecessary_skill_calls(frame: dict[str, Any], plan: dict[str, Any], allowed_attrs: set[str]) -> list[str]:
    required_fact_ids = {
        item.get("fact_id") or item.get("fact_requirement_id") or item.get("fact_id")
        for item in plan.get("facts") or []
        if item.get("priority") == "required" and item.get("attribute") in allowed_attrs
    }
    rows = []
    for call in plan.get("skill_calls") or []:
        covers = set(call.get("covers") or [])
        if call.get("covers_required_count", 0) == 0:
            rows.append(str(call.get("skill_id")))
        elif required_fact_ids and not covers.intersection(required_fact_ids) and call.get("covers_optional_count", 0) > 0:
            rows.append(str(call.get("skill_id")))
    return [item for item in unique(rows) if item and item != "None"]


def has_missing(skill_calls: list[dict[str, Any]], param: str) -> bool:
    return any(param in (item.get("missing_params") or []) for item in skill_calls)


def is_invalid_or_compliance(question: str) -> bool:
    return bool(
        re.search(
            r"明天|未来|2099|近100年|必须保证|稳赚不赔|一定会涨|预测明天|ABCDEF|不存在|(?<!\d)00309(?!\d)|(?<!\d)003095999(?!\d)",
            question,
        )
    )


def has_rejection_or_risk_issue(plan: dict[str, Any]) -> bool:
    text = json.dumps(plan.get("issues") or [], ensure_ascii=False)
    return bool(re.search(r"非法|未来|合规|预测|保证|风险|不支持|invalid|future|compliance", text, re.IGNORECASE))


def evaluation_text(
    grade: str,
    coverage: dict[str, Any],
    execution: dict[str, Any],
    fact_count: int,
    skill_count: int,
    has_diff: bool,
    answerability_issues: list[str],
    relation_issues: list[str],
    redundancy_issues: list[str],
    abnormal_issues: list[str],
    evaluation_outcome: str,
    expected_terminal_behavior: str,
    true_planner_error: bool,
    failure_cause: str,
    oag_reasoning_error: bool,
    evaluator_reason_zh: str,
) -> str:
    base = (
        f"{grade}：规划生成 {fact_count} 个事实需求、{skill_count} 个 Skill 调用；"
        f"覆盖状态为 {coverage.get('coverage_status')}，执行状态为 {execution.get('execution_status')}。"
    )
    parts = []
    parts.append(
        f"终态判断：期望 {expected_terminal_behavior}，实际 {evaluation_outcome}；"
        f"根因为 {failure_cause}；"
        f"{'存在 OAG 推理错误' if oag_reasoning_error else '未计为 OAG 推理错误'}。{evaluator_reason_zh}"
    )
    parts.append("回答完整性：" + ("存在缺口：" + "；".join(answerability_issues[:3]) if answerability_issues else "基本具备回答所需事实。"))
    parts.append("关系正确性：" + ("存在问题：" + "；".join(relation_issues[:3]) if relation_issues else "未发现显式关系错误。"))
    parts.append("冗余控制：" + ("存在冗余：" + "；".join(redundancy_issues[:3]) if redundancy_issues else "未发现明显不必要输出。"))
    if abnormal_issues:
        parts.append("异常/缺参处理：" + "；".join(abnormal_issues[:3]))
    if has_diff:
        parts.append("一致性：Java 与 Python 输出存在差异，需要优先定位。")
    return base + " " + " ".join(parts)


def summarize_plan(plan: dict[str, Any]) -> str:
    task = plan.get("task") or {}
    coverage = plan.get("coverage") or {}
    execution = plan.get("execution") or {}
    skills = [item.get("skill_id") for item in plan.get("skill_calls") or [] if item.get("skill_id")]
    facts = [
        item.get("attribute") or item.get("fact_type")
        for item in plan.get("facts") or []
        if item.get("attribute") or item.get("fact_type")
    ]
    issues = [item.get("code") for item in plan.get("issues") or [] if item.get("code")]
    return (
        f"status={plan.get('status')}; task={task.get('task_type')}/{task.get('intent')}; "
        f"coverage={coverage.get('coverage_status')}({coverage.get('covered_required_fact_count')}/{coverage.get('required_fact_count')}); "
        f"execution={execution.get('execution_status')}; skills={','.join(skills[:6]) or '-'}; "
        f"facts={','.join(facts[:10]) or '-'}; issues={','.join(issues[:6]) or '-'}"
    )


def build_summary(
    cases: list[dict[str, Any]],
    summary_counter: dict[str, Counter[str]],
    issue_counter: Counter[str],
    exact_match_count: int,
    outcome_counter: Counter[str],
    category_outcome_counter: dict[str, Counter[str]],
    failure_cause_counter: Counter[str],
    category_failure_cause_counter: dict[str, Counter[str]],
) -> dict[str, Any]:
    categories = []
    total_counter = Counter()
    outcome_keys = [
        "answerable_ready",
        "correct_blocked_missing_params",
        "correct_unsupported_data_source",
        "correct_unsupported_period",
        "correct_rejected_guardrail",
        "true_planner_error",
        "ambiguous_needs_review",
    ]
    failure_cause_keys = [
        "none",
        "system_limitation",
        "input_missing",
        "guardrail_policy",
        "oag_reasoning_error",
        "needs_review",
    ]
    for category in sorted(summary_counter):
        counter = summary_counter[category]
        outcomes = category_outcome_counter.get(category, Counter())
        causes = category_failure_cause_counter.get(category, Counter())
        total = sum(counter.values())
        total_counter.update(counter)
        categories.append(
            {
                "category": category,
                "total": total,
                "A": counter.get("A", 0),
                "B": counter.get("B", 0),
                "C": counter.get("C", 0),
                "D": counter.get("D", 0),
                "ERR": counter.get("ERR", 0),
                **{key: outcomes.get(key, 0) for key in outcome_keys},
                **{f"cause_{key}": causes.get(key, 0) for key in failure_cause_keys},
            }
        )
    total = len(cases)
    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_cases": total,
        "exact_match_count": exact_match_count,
        "exact_match_rate": round(exact_match_count / total, 4) if total else 0,
        "grade_totals": {grade: total_counter.get(grade, 0) for grade in ["A", "B", "C", "D", "ERR"]},
        "outcome_totals": {key: outcome_counter.get(key, 0) for key in outcome_keys},
        "failure_cause_totals": {key: failure_cause_counter.get(key, 0) for key in failure_cause_keys},
        "categories": categories,
        "top_issues": [{"issue": issue, "count": count} for issue, count in issue_counter.most_common(20)],
    }


def cell_json(value: Any, sidecar_name: str) -> str:
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if len(text) <= MAX_EXCEL_CELL_CHARS:
        return text
    sidecar = OUTPUT_DIR / "sidecars" / sidecar_name
    sidecar.parent.mkdir(parents=True, exist_ok=True)
    sidecar.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return json.dumps(
        {
            "truncated": True,
            "sidecar": str(sidecar),
            "summary": compact(value, limit=3000),
        },
        ensure_ascii=False,
    )


def compact(value: Any, limit: int = 600) -> str:
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return text if len(text) <= limit else text[: limit - 3] + "..."


def unique(values: list[Any]) -> list[Any]:
    result = []
    seen = set()
    for value in values:
        if value in (None, ""):
            continue
        key = json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else value
        if key in seen:
            continue
        seen.add(key)
        result.append(value)
    return result


def unique_dicts(values: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return unique(values)


if __name__ == "__main__":
    main()
